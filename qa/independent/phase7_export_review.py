"""Parse real browser downloads; publish counts/hashes only, never unfiltered report contents."""
import argparse,csv,io,json,re,hashlib
from pathlib import Path
import openpyxl,fitz
p=argparse.ArgumentParser()
for n in ['input','judgments','sha','output']:p.add_argument('--'+n,required=True)
a=p.parse_args();out=Path(a.output);out.mkdir(exist_ok=True)
if(out/'results.json').exists():raise RuntimeError('Preserve evidence')
results=[]
def check(id,expected,actual,ok):
 results.append(dict(id=id,story_ids=['LD-034'],priority='P1',preconditions=['Exact attested UAT SHA','Actual authenticated browser downloads; private files outside repository'],test_data={'source':'Five reports, English and Arabic; populated owned court for judgments'},steps=['Read downloaded CSV/XLSX/PDF with independent libraries','Compare CSV and workbook rows to live report API total','Check literal cell types, RTL, embedded PDF fonts and neutral text preservation'],expected_results=expected,actual=actual,status='PASS'if ok else'FAIL'))
for name in ['cases','judgments','hearings','experts','workload']:
 root=Path(a.judgments if name=='judgments'else a.input);meta=json.loads((root/'results.json').read_text())['artifacts']
 for lang in ['en','ar']:
  csvfile=root/f'{name}-{lang}.csv';rows=list(csv.reader(io.StringIO(csvfile.read_text(encoding='utf-8-sig'))));api=next(x['api_rows']for x in meta if x['report']==name and x['lang']==lang and x['format']=='csv')
  w=openpyxl.load_workbook(root/f'{name}-{lang}.xlsx').active;stream=list(openpyxl.load_workbook(root/f'{name}-{lang}.xlsx',read_only=True).active.values);formula=sum(c.data_type=='f'for row in w for c in row)
  check('P7-EXPORT-ROWS-'+name+'-'+lang,'CSV and normal/streaming XLSX contain exactly API count+header; no formulas; RTL iff Arabic',{'api_rows':api,'csv_rows':len(rows)-1,'xlsx_rows':w.max_row-1,'stream_rows':len(stream)-1,'formula_cells':formula,'rtl':bool(w.sheet_view.rightToLeft)},len(rows)==w.max_row==len(stream)==api+1 and formula==0 and bool(w.sheet_view.rightToLeft)==(lang=='ar'))
  pdf=fitz.open(root/f'{name}-{lang}.pdf');text=''.join(page.get_text()for page in pdf);flat=''.join(text.split());embedded=any(bool(pdf.extract_font(font[0])[3])for page in pdf for font in page.get_fonts(full=True));tokens=[r[0]for r in rows[1:]if re.fullmatch(r'[A-Za-z0-9_.-]{8,}',r[0])];found=sum(t in flat for t in tokens)
  check('P7-EXPORT-PDF-'+name+'-'+lang,'Parseable PDF with embedded fonts; every eligible neutral first-column identifier preserved',{'pages':len(pdf),'embedded_font':embedded,'neutral_tokens_checked':len(tokens),'neutral_tokens_found':found,'pdf_sha256':hashlib.sha256((root/f'{name}-{lang}.pdf').read_bytes()).hexdigest()},len(pdf)>0 and embedded and found==len(tokens));pdf.close()
(out/'results.json').write_text(json.dumps({'tested_sha':a.sha,'scope':'Independent parser oracle on all five browser exports; no private report data published','counts':{s:sum(r['status']==s for r in results)for s in ['PASS','FAIL','BLOCKED','NOT_RUN']},'results':results},indent=2));print(json.dumps({s:sum(r['status']==s for r in results)for s in ['PASS','FAIL','BLOCKED','NOT_RUN']}))
