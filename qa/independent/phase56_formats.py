"""Independent XLSX/PDF artifacts from real report endpoints; no developer assertions reused."""
import argparse,datetime,io,json,re,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl,fitz
from phase56_operations import OperationReview
from phase234_review import Client

class FormatReview(OperationReview):
    def req(self,id,method,path,body=None,status=200,client=None,rule=None):return self.request(id,'LD-034',client or self.admin,method,path,body,status,rule)
    def oracle(self,id,steps,expected,actual,ok):self.check(id,'LD-034',steps,expected,actual,ok)
    def formats(self):
        c=self.case('QA قضية عربية (123/2026)');p=self.proceeding(c);self.browser_case=c;self.browser_proceeding=p
        self.req('FORMAT-JUDGMENT-FIXTURE','POST','/api/v1/proceedings/'+p['id']+'/judgment',dict(judgmentDate='2026-01-03',outcome='IN_FAVOUR',final=True),201)
        formula=self.case('=1+1');self.case('+SUM(1,1)');self.case('-1+1');self.case('@SUM(1,1)');self.case('QA comma, "quoted" عربي')
        base='/api/v1/reports/cases/export';artifacts=self.output/'artifacts';artifacts.mkdir()
        for fmt in ['xlsx','pdf']:
            self.req('FORMAT-ANON-'+fmt,'GET',base+'?format='+fmt,status=401,client=Client(self.appport))
            self.req('FORMAT-DENIED-'+fmt,'GET',base+'?format='+fmt,status=403,client=self.none)
            for lang in ['en','ar']:
                s,data,h=self.raw(self.admin,'GET',base+'?format='+fmt+'&lang='+lang+'&sort=referenceNo,asc');file=artifacts/(lang+'.'+fmt)
                if s==200:file.write_bytes(data)
                self.oracle('FORMAT-RESPONSE-'+fmt+'-'+lang,'Export six real case fixtures','200 attachment,nosniff,no-store,correct MIME',dict(status=s,bytes=len(data),type=h.get('content-type'),attachment=h.get('content-disposition'),cache=h.get('cache-control')),s==200 and 'attachment' in h.get('content-disposition','') and h.get('x-content-type-options')=='nosniff' and 'no-store' in h.get('cache-control','') and ('application/pdf' in h.get('content-type','') if fmt=='pdf' else 'spreadsheetml.sheet' in h.get('content-type','')))
                if s!=200:continue
                if fmt=='xlsx':
                    wb=openpyxl.load_workbook(io.BytesIO(data),data_only=False);ws=wb.active;values=list(ws.values)
                    with zipfile.ZipFile(io.BytesIO(data)) as archive:
                        sheet=ET.fromstring(archive.read('xl/worksheets/sheet1.xml'));dimension=sheet.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension').get('ref')
                    stream=openpyxl.load_workbook(io.BytesIO(data),read_only=True).active
                    streamed=sum(1 for _ in stream.rows)
                    self.oracle('XLSX-DIMENSION-STREAMING-'+lang,'Compare OOXML dimension, actual rows and independent streaming reader','All7 rows available; dimension bounds actual data',dict(dimension=dimension,normal_rows=ws.max_row,streaming_rows=streamed),streamed==ws.max_row==7)
                    self.oracle('XLSX-RECORDS-'+lang,'Open workbook independently and compare fixtures','Header+6rows,9columns; Arabic title exact; no formulas',dict(rows=ws.max_row,columns=ws.max_column,titles=[r[1] for r in values[1:]],formula_count=sum(cell.data_type=='f' for row in ws for cell in row)),ws.max_row==7 and ws.max_column==9 and values[1][1]==c['title'] and all(cell.data_type!='f' for row in ws for cell in row))
                    self.oracle('XLSX-TYPES-RTL-'+lang,'Inspect dates/count cells and sheet view','Actual dates/numbers; RTL iff Arabic',dict(date_type=type(ws['F2'].value).__name__,count_type=type(ws['I2'].value).__name__,rtl=ws.sheet_view.rightToLeft),isinstance(ws['F2'].value,datetime.datetime) and isinstance(ws['I2'].value,(int,float)) and bool(ws.sheet_view.rightToLeft)==(lang=='ar'))
                    self.oracle('XLSX-FORMULA-TEXT-'+lang,'Inspect formula-shaped titles as cells','Exact literal string values stored as strings',dict(cells=[(ws.cell(r,2).value,ws.cell(r,2).data_type) for r in range(3,7)]),all(ws.cell(r,2).data_type=='s' for r in range(3,7)) and ws['B3'].value=='=1+1')
                else:
                    pdf=fitz.open(stream=data,filetype='pdf');page=pdf[0];fonts=page.get_fonts(full=True);text=''.join(p.get_text() for p in pdf)
                    page.get_pixmap(matrix=fitz.Matrix(1.8,1.8)).save(str(artifacts/(lang+'-page1.png')))
                    (artifacts/(lang+'-text.txt')).write_text(text)
                    self.oracle('PDF-STRUCTURE-'+lang,'Parse exported PDF and render first page','Readable landscape PDF; embedded font; no active script',dict(pages=len(pdf),width=page.rect.width,height=page.rect.height,font_names=[f[3] for f in fonts],has_embedded_font=any(bool(pdf.extract_font(f[0])[3]) for f in fonts)),len(pdf)>=1 and page.rect.width>page.rect.height and any(bool(pdf.extract_font(f[0])[3]) for f in fonts) and b'/JavaScript' not in data)
                    self.oracle('PDF-IDENTIFIERS-'+lang,'Inspect extracted text for Latin case reference and numeric expression','Known reference present without lost data',dict(reference_present=c['referenceNo'] in text,formula_present='=1+1' in text),c['referenceNo'] in text and '=1+1' in text)
                    pdf.close()
        for fmt in ['xlsx','pdf']:
            s,data,_=self.raw(self.admin,'GET','/api/v1/reports/judgments/export?format='+fmt+'&lang=ar&from=2026-01-03&to=2026-01-03&final=true')
            (artifacts/('judgments-ar.'+fmt)).write_bytes(data)
            if fmt=='xlsx':count=openpyxl.load_workbook(io.BytesIO(data)).active.max_row-1
            else:
                doc=fitz.open(stream=data,filetype='pdf');count=int(c['referenceNo'] in ''.join(p.get_text() for p in doc));doc[0].get_pixmap(matrix=fitz.Matrix(1.8,1.8)).save(str(artifacts/'judgments-ar-page1.png'));doc.close()
            self.oracle('FORMAT-JUDGMENT-FILTER-'+fmt,'Export exact day and finality','Exactly one fixture judgment',dict(status=s,count=count),s==200 and count==1)
        # Repeated page headers and wrapped Arabic/Latin values at ordinary multi-page volume.
        for i in range(60):self.case('QA ملف عربي طويل '+str(i)+' / رقم 123/2026 '+('تفاصيل ' * 8))
        s,data,_=self.raw(self.admin,'GET',base+'?format=pdf&lang=ar&sort=referenceNo,asc');(artifacts/'multipage-ar.pdf').write_bytes(data)
        doc=fitz.open(stream=data,filetype='pdf');texts=[page.get_text() for page in doc];first_header=texts[0].splitlines()[1:4]
        self.oracle('PDF-MULTIPAGE','Render66 cases with long bilingual titles','Multiple nonempty pages; representative header repetition reviewed visually in separate evidence',dict(status=s,pages=len(doc),nonempty_pages=[bool(t.strip()) for t in texts]),s==200 and len(doc)>1 and all(t.strip() for t in texts))
        doc[-1].get_pixmap(matrix=fitz.Matrix(1.8,1.8)).save(str(artifacts/'multipage-ar-last.png'));doc.close()
        self.sql("insert into legal_case(reference_no,title,case_type,opened_on,created_by,updated_by) select 'QA-FORMAT-CAP-'||g,'QA cap','CIVIL','2026-01-01',u.id,u.id from generate_series(1,9934) g cross join app_user u where u.username='qaadmin'")
        s,data,_=self.raw(self.admin,'GET',base+'?format=xlsx');(artifacts/'cap-10000.xlsx').write_bytes(data);rows=openpyxl.load_workbook(io.BytesIO(data),read_only=False).active.max_row
        self.oracle('XLSX-LIMIT-10000','Export exactly10000 cases','200 and10000data rows',dict(status=s,rows=rows),s==200 and rows==10001)
        s,data,_=self.raw(self.admin,'GET',base+'?format=pdf');(artifacts/'cap-10000.pdf').write_bytes(data)
        capdoc=fitz.open(stream=data,filetype='pdf');captext='\n'.join(p.get_text() for p in capdoc)
        # The narrow reference cell wraps digits; join only within that fixture's
        # reference, bounded by its known title, rather than counting first digits.
        capids=[int(re.sub(r'\s+','',n)) for n in re.findall(r'QA-FORMAT-CAP-([\d\s]+)(?=QA cap)',captext)]
        found=len(set(capids) & set(range(1,9935)))
        self.oracle('PDF-LIMIT-10000','Export exactly10000 records and inspect generated PDF','200; PDF contains all cap fixture reference tokens',dict(status=s,pages=len(capdoc),cap_refs_present=found),s==200 and found==9934);capdoc.close()
        self.case('QA above cap')
        for fmt in ['xlsx','pdf']:self.req('FORMAT-LIMIT-10001-'+fmt,'GET',base+'?format='+fmt,status=422,rule='EXPORT_TOO_LARGE')
        self.sql("create function qa_export_audit_fail() returns trigger language plpgsql as $$ begin if new.action='REPORT_EXPORT' then raise exception 'QA injected audit failure'; end if; return new; end $$; create trigger qa_export_audit_fail before insert on audit_event for each row execute function qa_export_audit_fail()")
        before=self.sql("select count(*) from audit_event where action='REPORT_EXPORT'").stdout
        status,data,headers=self.raw(self.admin,'GET',base+'?format=xlsx&q=QA56')
        after=self.sql("select count(*) from audit_event where action='REPORT_EXPORT'").stdout
        self.sql('drop trigger qa_export_audit_fail on audit_event; drop function qa_export_audit_fail()')
        self.oracle('EXPORT-AUDIT-FAIL-CLOSED','Inject audit failure in disposable DB then request valid small XLSX','500; no attachment body or successful audit row',dict(status=status,attachment=headers.get('content-disposition'),audit_unchanged=before==after),status==500 and not data.startswith(b'PK') and 'attachment' not in headers.get('content-disposition','') and before==after)
        # Explicit operational blocker: local jar success cannot stand in for deployment-image execution.
        self.results.append(dict(id='PDF-ACTUAL-DEPLOYMENT-CONTAINER',story_ids=['LD-034','LD-006'],priority='P1',preconditions=['Approved isolated access to actual deployment image/container'],test_data={'formats':['pdf'],'lang':'ar'},steps=['Generate Arabic PDF through exact deployment-container API','Render artifact and compare to jar output'],expected_results='Valid embedded/shaped PDF inside production-equivalent image',status='BLOCKED',actual='Local Docker socket permission denied; existing sudo -n requires password. No daemon/group/infrastructure changes attempted.'));self.save()

if __name__=='__main__':
 p=argparse.ArgumentParser()
 for name in ['candidate','sha','output']:p.add_argument('--'+name,required=True)
 p.add_argument('--kind',default='formats');p.add_argument('--browser',action='store_true')
 raise SystemExit(FormatReview(p.parse_args()).execute())
