"""Independent parser checks for own-prefix exports obtained from authenticated UAT."""
import argparse,csv,io,json,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl,fitz
p=argparse.ArgumentParser()
for n in ['input','sha','output']:p.add_argument('--'+n,required=True)
a=p.parse_args();root=Path(a.input);out=Path(a.output);out.mkdir(exist_ok=True)
if (out/'results.json').exists():raise RuntimeError('Preserve evidence')
results=[]
def check(id,expected,actual,ok):
    results.append({'id':id,'story_ids':['LD-034'],'priority':'P1','preconditions':['Approved-account UAT export, exact runtime SHA, synthetic prefix filter'],'test_data':{'artifact_directory':root.name},'steps':['Download authenticated filtered export from deployed backend','Parse with independent CSV/openpyxl/PyMuPDF consumer','Compare row counts, neutral identifiers, embedded fonts and RTL'], 'expected_results':expected,'actual':actual,'status':'PASS'if ok else 'FAIL'})
for lang in ['en','ar']:
    csvrows=list(csv.reader(io.StringIO((root/f'owned-cases-{lang}.csv').read_text(encoding='utf-8-sig'))))
    f=root/f'owned-cases-{lang}.xlsx';normal=openpyxl.load_workbook(f).active;streamed=list(openpyxl.load_workbook(f,read_only=True).active.values)
    with zipfile.ZipFile(f)as z:dimension=ET.fromstring(z.read('xl/worksheets/sheet1.xml')).find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension').get('ref')
    check('UAT-QAD035-'+lang,'Normal and streaming XLSX consumers see every CSV row',{'csv_rows':len(csvrows),'normal_rows':normal.max_row,'streaming_rows':len(streamed),'dimension':dimension},len(csvrows)==normal.max_row==len(streamed))
    check('UAT-XLSX-LITERAL-RTL-'+lang,'Literal =1+1 string; zero formulas; RTL iff Arabic',{'formula_count':sum(c.data_type=='f'for row in normal for c in row),'neutral_expression_present':any(c.value=='=1+1'and c.data_type=='s'for row in normal for c in row),'rtl':bool(normal.sheet_view.rightToLeft)},all(c.data_type!='f'for row in normal for c in row)and any(c.value=='=1+1'and c.data_type=='s'for row in normal for c in row)and bool(normal.sheet_view.rightToLeft)==(lang=='ar'))
    check('UAT-CSV-FORMULA-'+lang,'Formula-shaped title is quoted literal',{'neutralized':any("'=1+1"in row for row in csvrows)},any("'=1+1"in row for row in csvrows))
    pdf=fitz.open(root/f'owned-cases-{lang}.pdf');text=''.join(page.get_text()for page in pdf);embedded=any(bool(pdf.extract_font(font[0])[3])for page in pdf for font in page.get_fonts(full=True))
    pdf[0].get_pixmap(matrix=fitz.Matrix(1.6,1.6)).save(str(out/f'uat-{lang}-page1.png'))
    check('UAT-QAD036-'+lang,'Neutral expression remains =1+1; synthetic case reference retained',{'neutral_expression_present':'=1+1'in text,'known_reference_present':csvrows[1][0]in ''.join(text.split())},'=1+1'in text and csvrows[1][0]in ''.join(text.split()))
    check('UAT-PDF-DEPLOYED-'+lang,'Parseable landscape PDF with embedded font generated through deployed API',{'pages':len(pdf),'embedded_font':embedded,'landscape':pdf[0].rect.width>pdf[0].rect.height},len(pdf)>0 and embedded and pdf[0].rect.width>pdf[0].rect.height);pdf.close()
(out/'results.json').write_text(json.dumps({'tested_sha':a.sha,'scope':'Independent parsing/rendering of actual authenticated UAT artifacts; deployment container provenance linked separately','counts':{s:sum(r['status']==s for r in results)for s in ['PASS','FAIL','BLOCKED','NOT_RUN']},'results':results},indent=2))
print(json.dumps({s:sum(r['status']==s for r in results)for s in ['PASS','FAIL','BLOCKED','NOT_RUN']}))
